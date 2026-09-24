import json
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.export_service import _excel_text, create_final_exports, resolve_export_path


def test_excel_text_blocks_formula_prefixes() -> None:
    assert _excel_text("=SUM(1,2)") == "'=SUM(1,2)"


def test_export_path_rejects_unexpected_filename_without_scanning_storage() -> None:
    assert resolve_export_path("../ket-qua.xlsx") is None
    assert resolve_export_path("ket-qua.xlsx") is None


def test_final_export_creates_matching_json_and_excel(monkeypatch, tmp_path) -> None:
    reading = SimpleNamespace(final_customer_id="PN2.001", reading_value=Decimal("63751"))
    ai_result = SimpleNamespace(final_confidence=0.8208)
    customer = SimpleNamespace(
        customer_code="PN2.001",
        full_name="Nguyễn Văn A",
        address="Khu 1",
        electricity_route="Tuyến 02",
        meter_serial="CT-2026-001",
        initial_reading=13836,
        usage_purpose="SINH_HOAT",
    )
    db = SimpleNamespace(execute=lambda _statement: [(reading, ai_result, customer)])
    monkeypatch.setattr(
        "app.services.export_service.get_settings",
        lambda: SimpleNamespace(storage_root=tmp_path),
    )

    bundle = create_final_exports(db, month=9, year=2026)

    payload = json.loads(bundle.json_path.read_text(encoding="utf-8"))
    assert payload == [
        {
            "ma_khach_hang": "PN2.001",
            "ho_ten": "Nguyễn Văn A",
            "dia_chi": "Khu 1",
            "tuyen_dien": "Tuyến 02",
            "so_seri_cong_to": "CT-2026-001",
            "chi_so_khoi_tao": 13836,
            "muc_dich_su_dung": "SINH_HOAT",
            "chi_so_moi": 63751,
            "do_tin_cay": 0.8208,
            "ket_qua_doi_chieu": "KHOP",
        }
    ]
    workbook = load_workbook(bundle.excel_path, data_only=True)
    sheet = workbook["Chỉ số 09-2026"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Mã khách hàng", "Số điện", "Độ tin cậy")
    assert rows[1] == ("PN2.001", 63751, 0.8208)
    assert sheet.auto_filter.ref == "A1:C2"
    assert "-ky-2026-09" in bundle.excel_name
    workbook.close()
