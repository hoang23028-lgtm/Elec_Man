import json
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models.ai_result import AiResult
from app.models.customer import Customer
from app.models.meter_reading import MeterReading

EXPORT_FILENAME_PATTERN = re.compile(
    r"^bao-cao-doi-chieu-(\d{4})(\d{2})\d{2}T\d{6}Z-[0-9a-f]{8}\.(xlsx|json)$"
)
LEGACY_EXPORT_FILENAME_PATTERN = re.compile(
    r"^chi-so-da-xac-nhan-(\d{4})(\d{2})\d{2}T\d{6}Z-[0-9a-f]{8}\.xlsx$"
)


@dataclass(frozen=True)
class ExportBundle:
    excel_name: str
    excel_path: Path
    json_name: str
    json_path: Path
    row_count: int


def _excel_text(value: str | None) -> str:
    text = value or ""
    return f"'{text}" if text[:1] in {"=", "+", "-", "@"} else text


def resolve_export_path(filename: str) -> Path | None:
    match = EXPORT_FILENAME_PATTERN.fullmatch(filename)
    if match is None:
        match = LEGACY_EXPORT_FILENAME_PATTERN.fullmatch(filename)
    if match is None:
        return None
    root = (get_settings().storage_root / "exports").resolve()
    path = (root / match.group(1) / match.group(2) / filename).resolve()
    return path if path.is_relative_to(root) and path.is_file() else None


def _report_row(
    reading: MeterReading, ai_result: AiResult, customer: Customer | None
) -> dict[str, str | int | float | None]:
    reading_value = reading.reading_value
    return {
        "ma_khach_hang": customer.customer_code if customer else reading.final_customer_id,
        "ho_ten": customer.full_name if customer else None,
        "dia_chi": customer.address if customer else None,
        "tuyen_dien": customer.electricity_route if customer else None,
        "so_seri_cong_to": customer.meter_serial if customer else None,
        "chi_so_khoi_tao": customer.initial_reading if customer else None,
        "muc_dich_su_dung": customer.usage_purpose if customer else None,
        "chi_so_moi": int(reading_value) if isinstance(reading_value, Decimal) else reading_value,
        "do_tin_cay": round(ai_result.final_confidence, 6),
        "ket_qua_doi_chieu": "KHOP" if customer else "KHONG_TIM_THAY",
    }


def create_final_exports(db: Session) -> ExportBundle:
    statement = (
        select(MeterReading, AiResult, Customer)
        .join(AiResult, AiResult.id == MeterReading.ai_result_id)
        .outerjoin(Customer, Customer.id == MeterReading.matched_customer_id)
        .where(MeterReading.review_status == "CONFIRMED")
        .order_by(MeterReading.created_at)
        .execution_options(yield_per=1000)
    )
    now = datetime.now(UTC)
    token = uuid4().hex[:8]
    stem = f"bao-cao-doi-chieu-{now:%Y%m%dT%H%M%SZ}-{token}"
    directory = get_settings().storage_root / "exports" / f"{now:%Y}" / f"{now:%m}"
    directory.mkdir(parents=True, exist_ok=True)
    excel_name = f"{stem}.xlsx"
    json_name = f"{stem}.json"
    excel_path = directory / excel_name
    json_path = directory / json_name
    excel_temporary = excel_path.with_suffix(".xlsx.tmp")
    json_temporary = json_path.with_suffix(".json.tmp")

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Báo cáo đối chiếu")
    sheet.freeze_panes = "A2"
    widths = {
        "A": 22,
        "B": 28,
        "C": 42,
        "D": 34,
        "E": 20,
        "F": 18,
        "G": 20,
        "H": 16,
        "I": 16,
        "J": 20,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    headers = [
        "Mã khách hàng",
        "Họ tên",
        "Địa chỉ",
        "Tuyến điện",
        "Số serial công tơ",
        "Chỉ số khởi tạo",
        "Mục đích sử dụng",
        "Chỉ số mới",
        "Độ tin cậy",
        "Kết quả đối chiếu",
    ]
    header_cells = [WriteOnlyCell(sheet, value=value) for value in headers]
    for cell in header_cells:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B7285")
    sheet.append(header_cells)

    row_count = 0
    try:
        with json_temporary.open("w", encoding="utf-8", newline="\n") as output:
            output.write("[\n")
            for reading, ai_result, customer in db.execute(statement):
                report = _report_row(reading, ai_result, customer)
                if row_count:
                    output.write(",\n")
                output.write("  ")
                json.dump(report, output, ensure_ascii=False)
                confidence_cell = WriteOnlyCell(sheet, value=report["do_tin_cay"])
                confidence_cell.number_format = "0.00%"
                sheet.append(
                    [
                        _excel_text(report["ma_khach_hang"]),
                        _excel_text(report["ho_ten"]),
                        _excel_text(report["dia_chi"]),
                        _excel_text(report["tuyen_dien"]),
                        _excel_text(report["so_seri_cong_to"]),
                        report["chi_so_khoi_tao"],
                        _excel_text(report["muc_dich_su_dung"]),
                        report["chi_so_moi"],
                        confidence_cell,
                        report["ket_qua_doi_chieu"],
                    ]
                )
                row_count += 1
            output.write("\n]\n")
        sheet.auto_filter.ref = f"A1:J{max(1, row_count + 1)}"
        workbook.save(excel_temporary)
        excel_temporary.replace(excel_path)
        json_temporary.replace(json_path)
    except Exception:
        excel_temporary.unlink(missing_ok=True)
        json_temporary.unlink(missing_ok=True)
        excel_path.unlink(missing_ok=True)
        json_path.unlink(missing_ok=True)
        raise
    return ExportBundle(
        excel_name=excel_name,
        excel_path=excel_path,
        json_name=json_name,
        json_path=json_path,
        row_count=row_count,
    )
