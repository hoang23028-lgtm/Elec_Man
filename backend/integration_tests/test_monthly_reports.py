from io import BytesIO

from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import event

from app.core.database import SessionLocal, engine
from app.services.ai_result_service import store_immutable_ai_result
from app.services.job_service import claim_next_job, mark_job_completed

API = "/api/v1"
FIELDS = (
    "ma_khach_hang",
    "ho_ten",
    "dia_chi",
    "tuyen_dien",
    "so_seri_cong_to",
    "chi_so_khoi_tao",
    "muc_dich_su_dung",
)


def ok(response, code=200):
    assert response.status_code == code, response.text
    return response.json()


def test_monthly_dashboard_and_exports_follow_confirmed_history(login):
    client = login("ADMIN")
    ok(
        client.post(
            API + "/customers",
            json={
                "customer_code": "REPORT.001",
                "full_name": "Khách kiểm thử báo cáo",
                "address": "Dữ liệu thử",
                "electricity_route": "TEST",
                "meter_serial": "REPORT001",
                "initial_reading": 100,
                "usage_purpose": "SINH_HOAT",
            },
        ),
        201,
    )
    batch = ok(client.post(API + "/batches", json={"original_folder_name": "reports"}), 201)
    image_ids = []
    for color in ("purple", "yellow", "orange"):
        stream = BytesIO()
        Image.new("RGB", (640, 480), color).save(stream, format="PNG")
        uploaded = ok(
            client.post(
                API + f"/batches/{batch['id']}/images",
                files={
                    "file": (color + ".png", stream.getvalue(), "image/png"),
                },
            ),
            201,
        )
        image_ids.append(uploaded["id"])
    ok(client.post(API + f"/batches/{batch['id']}/start"))
    with SessionLocal() as db:
        for _ in image_ids:
            job = claim_next_job(db, "report-worker")
            result = {
                "model_version": "integration-fixture",
                "customer_id_ai": None,
                "meter_reading_ai": "00120",
                "customer_confidence": 0,
                "meter_confidence": 0.8,
                "detection_confidence": 0.8,
                "image_quality_score": 0.8,
                "final_confidence": 0.8,
                "status": "REVIEW_REQUIRED",
                "processing_time_ms": 1,
            }
            store_immutable_ai_result(db, job.image_id, result)
            db.commit()
            mark_job_completed(db, job.id, result)

    def review(index, action, value=None, month=None):
        row = next(r for r in ok(client.get(API + "/results")) if r["image_id"] == image_ids[index])
        return ok(
            client.put(
                API + f"/results/{image_ids[index]}/review",
                json={
                    "expected_revision": row["revision"],
                    "action": action,
                    "final_customer_id": "REPORT.001",
                    "final_meter_reading": value,
                    "reading_month": f"2026-{month:02d}-01" if month else None,
                },
            )
        )

    def dashboard(month=None):
        query = "customer_id=REPORT.001&year=2026" + (f"&month={month}" if month else "")
        data = ok(client.get(API + "/dashboard/billing?" + query))
        assert data["total"] == len(data["records"])
        assert data["summary"]["total_records"] == len(data["records"])
        assert data["summary"]["total_consumption_kwh"] == sum(
            r["consumption_kwh"] or 0 for r in data["records"]
        )
        for field in ("energy_charge_before_vat_vnd", "vat_amount_vnd", "estimated_amount_vnd"):
            assert data["summary"][field] == sum(r[field] or 0 for r in data["records"])
            assert data["summary"][field] == sum(r[field] for r in data["trend"])
        return data

    def export(month, expected):
        bundle = ok(client.post(API + f"/exports/final?month={month}&year=2026"))
        json_rows = ok(client.get(bundle["json_download_url"]))
        assert bundle["exported_rows"] == len(json_rows) == (0 if expected is None else 1)
        workbook = load_workbook(
            BytesIO(client.get(bundle["excel_download_url"]).content),
            read_only=True,
            data_only=True,
        )
        try:
            sheet = workbook[f"Chỉ số {month:02d}-2026"]
            rows = list(sheet.iter_rows(values_only=True))
            assert rows[0] == FIELDS
            assert rows[1:] == [tuple(row[field] for field in FIELDS) for row in json_rows]
        finally:
            workbook.close()
        if expected is not None:
            assert tuple(json_rows[0]) == FIELDS
            assert json_rows[0]["ma_khach_hang"] == "REPORT.001"
            assert json_rows[0]["chi_so_khoi_tao"] == expected
            data = dashboard(month)
            assert int(data["records"][0]["current_reading"]) == expected
        return bundle

    assert dashboard()["total"] == 0
    export(1, None)  # Uploaded/OCR-completed images are not official monthly data.
    review(0, "CONFIRM", "00150", 1)
    review(1, "CONFIRM", "00180", 2)
    review(2, "CONFIRM", "00220", 3)
    assert [dashboard(month)["records"][0]["consumption_kwh"] for month in (1, 2, 3)] == [
        50,
        30,
        40,
    ]
    export(1, 150)
    old_bundle = export(2, 180)
    review(1, "CONFIRM", "00190", 2)
    assert dashboard(2)["records"][0]["consumption_kwh"] == 40
    assert dashboard(3)["records"][0]["consumption_kwh"] == 30
    export(2, 190)
    # Edit after the totals query, before the trend/detail queries. One response
    # must retain a single snapshot; the following request must see the edit.
    changed = False

    def edit_during_dashboard(conn, cursor, statement, parameters, context, executemany):
        nonlocal changed
        if not changed and "total_records" in statement and "confirmed_meter_history" in statement:
            changed = True
            review(1, "CONFIRM", "00200", 2)

    event.listen(engine, "after_cursor_execute", edit_during_dashboard)
    try:
        snapshot = dashboard(2)
        assert changed
        assert snapshot["records"][0]["current_reading"] == "190"
    finally:
        event.remove(engine, "after_cursor_execute", edit_during_dashboard)
    assert dashboard(2)["records"][0]["current_reading"] == "200"
    review(1, "CONFIRM", "00190", 2)
    # Previously downloaded files remain historical snapshots.
    assert ok(client.get(old_bundle["json_download_url"]))[0]["chi_so_khoi_tao"] == 180
    review(1, "REJECT")
    assert dashboard(2)["total"] == 0
    assert dashboard(3)["records"][0]["consumption_kwh"] is None
    export(2, None)
    review(1, "REOPEN")
    assert dashboard(2)["total"] == 0  # Reopening does not re-confirm official data.
    export(2, None)
    review(1, "CONFIRM", "00190", 2)
    assert dashboard(3)["records"][0]["consumption_kwh"] == 30
    export(2, 190)
    # A decreasing reading is visible, but cannot produce a negative bill.
    review(1, "CONFIRM", "00140", 2)
    assert dashboard(2)["records"][0]["consumption_kwh"] is None
    assert dashboard(2)["records"][0]["estimated_amount_vnd"] is None
    for index in range(3):
        review(index, "REJECT")
    assert dashboard()["total"] == 0
